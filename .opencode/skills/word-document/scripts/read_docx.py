from docx import Document
import sys
import json

def read_docx(file_path):
    doc = Document(file_path)
    content = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content.append({
                "type": "paragraph",
                "text": text,
                "style": para.style.name
            })

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            content.append({
                "type": "table",
                "data": table_data
            })

    return content

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python read_docx.py <file_path>")
        sys.exit(1)

    file_path = sys.argv[1]
    result = read_docx(file_path)
    print(json.dumps(result, ensure_ascii=False, indent=2))