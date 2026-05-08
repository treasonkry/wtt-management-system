from docx import Document
from docx.shared import Pt, RGBColor
import sys
import json

def replace_text(file_path, old_text, new_text):
    doc = Document(file_path)
    count = 0

    for para in doc.paragraphs:
        if old_text in para.text:
            para.text = para.text.replace(old_text, new_text)
            count += para.text.count(new_text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if old_text in cell.text:
                    cell.text = cell.text.replace(old_text, new_text)
                    count += 1

    doc.save(file_path)
    return count

def format_text(file_path, target_text, bold=None, italic=None, font_size=None, color=None):
    doc = Document(file_path)
    count = 0

    for para in doc.paragraphs:
        for run in para.runs:
            if target_text in run.text:
                if bold is not None:
                    run.bold = bold
                if italic is not None:
                    run.italic = italic
                if font_size is not None:
                    run.font.size = Pt(font_size)
                if color is not None:
                    run.font.color.rgb = RGBColor(*color)
                count += 1

    doc.save(file_path)
    return count

def add_text_after(file_path, search_text, new_text, style="Normal"):
    doc = Document(file_path)
    for para in doc.paragraphs:
        if search_text in para.text:
            new_para = para.insert_element_after_(para._element, new_text)
            break
    doc.save(file_path)

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python edit_docx.py <file_path> <action> <params_json>")
        print("Actions: replace, format")
        sys.exit(1)

    file_path = sys.argv[1]
    action = sys.argv[2]
    params = json.loads(sys.argv[3]) if len(sys.argv) > 3 else {}

    if action == "replace":
        count = replace_text(file_path, params.get("old_text", ""), params.get("new_text", ""))
        print(f"Replaced {count} occurrences")
    elif action == "format":
        count = format_text(
            file_path,
            params.get("target_text", ""),
            bold=params.get("bold"),
            italic=params.get("italic"),
            font_size=params.get("font_size"),
            color=params.get("color")
        )
        print(f"Formatted {count} runs")
    else:
        print(f"Unknown action: {action}")
        sys.exit(1)