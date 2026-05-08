from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys
import json

def create_docx(file_path, content_json):
    doc = Document()

    content = json.loads(content_json) if isinstance(content_json, str) else content_json

    if "title" in content:
        heading = doc.add_heading(content["title"], level=1)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    if "blocks" in content:
        for block in content["blocks"]:
            block_type = block.get("type", "paragraph")

            if block_type == "heading":
                doc.add_heading(block.get("text", ""), level=block.get("level", 1))

            elif block_type == "paragraph":
                doc.add_paragraph(block.get("text", ""))

            elif block_type == "table":
                data = block.get("data", [])
                if data:
                    table = doc.add_table(rows=len(data), cols=len(data[0]))
                    for i, row_data in enumerate(data):
                        for j, cell_text in enumerate(row_data):
                            table.rows[i].cells[j].text = str(cell_text)

            elif block_type == "image":
                img_path = block.get("path")
                if img_path:
                    doc.add_picture(img_path, width=Inches(block.get("width", 6)))

    doc.save(file_path)
    return file_path

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python create_docx.py <file_path> <content_json>")
        sys.exit(1)

    file_path = sys.argv[1]
    content_json = sys.argv[2]
    create_docx(file_path, content_json)
    print(f"Created: {file_path}")