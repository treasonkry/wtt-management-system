from openpyxl import load_workbook
import sys
import json

def read_excel(file_path, sheet_name=None):
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    data = []
    for row in ws.iter_rows(values_only=True):
        data.append(list(row))

    return {
        "sheet_name": ws.title,
        "data": data,
        "dimensions": ws.dimensions
    }

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python read_excel.py <file_path> [sheet_name]")
        sys.exit(1)

    file_path = sys.argv[1]
    sheet_name = sys.argv[2] if len(sys.argv) > 2 else None

    result = read_excel(file_path, sheet_name)
    print(json.dumps(result, ensure_ascii=False, indent=2))