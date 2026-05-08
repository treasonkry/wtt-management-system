from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import sys
import json

def set_cell(file_path, cell_ref, value):
    wb = load_workbook(file_path)
    ws = wb.active
    ws[cell_ref] = value
    wb.save(file_path)

def update_cells(file_path, updates_json):
    wb = load_workbook(file_path)
    ws = wb.active

    updates = json.loads(updates_json) if isinstance(updates_json, str) else updates_json
    for cell_ref, value in updates.items():
        ws[cell_ref] = value

    wb.save(file_path)

def append_row(file_path, row_data_json):
    wb = load_workbook(file_path)
    ws = wb.active

    row_data = json.loads(row_data_json) if isinstance(row_data_json, str) else row_data_json
    ws.append(row_data)

    wb.save(file_path)

def format_header(file_path, row_num=1, bg_color="CCE5FF"):
    wb = load_workbook(file_path)
    ws = wb.active

    header_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    header_font = Font(bold=True)

    for cell in ws[row_num]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    wb.save(file_path)

def add_formula(file_path, cell_ref, formula):
    wb = load_workbook(file_path)
    ws = wb.active
    ws[cell_ref] = formula
    wb.save(file_path)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python edit_excel.py <file_path> <action> <params>")
        print("Actions: set_cell, update_cells, append_row, format_header, add_formula")
        sys.exit(1)

    file_path = sys.argv[1]
    action = sys.argv[2]
    params = sys.argv[3] if len(sys.argv) > 3 else "{}"

    if action == "set_cell":
        p = json.loads(params)
        set_cell(file_path, p["cell"], p["value"])
        print(f"Set {p['cell']} = {p['value']}")
    elif action == "update_cells":
        update_cells(file_path, params)
        print("Updated cells")
    elif action == "append_row":
        append_row(file_path, params)
        print("Row appended")
    elif action == "format_header":
        p = json.loads(params)
        format_header(file_path, p.get("row", 1), p.get("bg_color", "CCE5FF"))
        print("Header formatted")
    elif action == "add_formula":
        p = json.loads(params)
        add_formula(file_path, p["cell"], p["formula"])
        print(f"Added formula to {p['cell']}")
    else:
        print(f"Unknown action: {action}")
        sys.exit(1)